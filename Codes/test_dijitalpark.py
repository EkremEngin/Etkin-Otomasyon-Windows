"""Dijitalpark taslakları ve sessiz kişi kaybı için portalsız regresyon kontrolleri."""
import datetime
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

import dgs_poc
import izin_data_v2
import izin_otomasyon


class DijitalparkWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def save(self, wb, name):
        path = Path(self.temp.name) / name
        wb.save(path)
        wb.close()
        return str(path)

    def dgs_book(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Page1"
        headers = {"A": "TCKimlikNo", "D": "Personel", "K": "Haftalık Çalışma Günü",
                   "L": "Haftalık Çalışma Saati", "N": "İstisnalı Çalışılan Toplam Süre",
                   "O": "Tam İstisna İçin Çalışılması Gereken Süre",
                   "P": "DGS Girilecek Süre", "R": "Yer Alınan Proje"}
        for col, value in headers.items():
            ws[f"{col}1"] = value
        ws["A3"] = "10*******18"
        ws["D3"] = "ÖRNEK KİŞİ"
        ws["K3"] = "5"
        ws["L3"] = "45"
        ws["N3"] = "29:49"
        ws["O3"] = "198:00"
        ws["P3"] = "=O3-N3"
        ws["R3"] = "Örnek proje"
        return wb

    def test_dgs_draft_formula_and_excel_schedule(self):
        path = self.save(self.dgs_book(), "draft.xlsx")
        people = dgs_poc.read_excel(path, "Eylül", "Dijitalpark")
        self.assertEqual(len(people), 1)
        person = next(iter(people.values()))
        self.assertEqual(person.planned_dgs_min, 168 * 60 + 11)
        self.assertFalse(person.manual_dgs)
        self.assertEqual(dgs_poc.person_schedule(person, {}), (540, False))

    def test_manual_shortening_never_enters_full_time(self):
        wb = self.dgs_book()
        wb.active["P3"] = "160:00"
        path = self.save(wb, "manual.xlsx")
        person = next(iter(dgs_poc.read_excel(path, "Eylül", "Dijitalpark").values()))
        self.assertTrue(person.manual_dgs)
        with self.assertRaisesRegex(dgs_poc.VerifyError, "manuel"):
            dgs_poc.process_one_person(None, person, True, {})

    def test_duplicate_person_and_wrong_park_stop_before_portal(self):
        wb = self.dgs_book()
        ws = wb.active
        for col in ("A", "D", "K", "L", "N", "O", "R"):
            ws[f"{col}4"] = ws[f"{col}3"].value
        ws["P4"] = "=O4-N4"
        path = self.save(wb, "duplicate.xlsx")
        with self.assertRaisesRegex(ValueError, "ikinci kez"):
            dgs_poc.read_excel(path, "Eylül", "Dijitalpark")
        with self.assertRaisesRegex(ValueError, "yalnız Dijitalpark"):
            dgs_poc.read_excel(path, "Eylül", "BV")

    def test_six_day_schedule_stays_closed(self):
        person = dgs_poc.PersonRow("ÖRNEK", "Proje", "00:00", "Ar-Ge", "Dijitalpark",
                                   weekly_days=6, weekly_hours=45)
        with self.assertRaisesRegex(dgs_poc.VerifyError, "CANLI TEST EDİLMEDİ"):
            dgs_poc.person_schedule(person, {})

    def test_old_dgs_layout_still_reads(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sayfa1"
        ws.append(["AD SOYAD", "TC Kimlik", "BÖLÜM", "Lokasyonu SGK",
                   "Ar-Ge/ Destek/ K.Dışı", "Eksik Puantaj (Saat)", "Proje Adı"])
        ws.append(["ÖRNEK KİŞİ", "10000000146", "Ar-Ge", "BV", "Ar-Ge", "09:00", "Proje"])
        path = self.save(wb, "old.xlsx")
        people = dgs_poc.read_excel(path, "Eylül", "BV")
        self.assertEqual(len(people), 1)
        self.assertIsNone(next(iter(people.values())).weekly_days)

    def test_izin_masked_tc_and_park_from_selection(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tabelle1"
        ws.append([])
        ws.append([])
        ws.append([])
        ws.append([None, None, "T.C.", "Ad-Soyad", "Tarih", "Gün Sayısı"])
        ws.append([None, None, "31**", "ÖRNEK BİR", datetime.datetime(2026, 9, 13), 1])
        ws.append([None, None, "31**", "ÖRNEK BİR", datetime.datetime(2026, 9, 14), 0.5])
        ws.append([None, None, "31**", "ÖRNEK İKİ", datetime.datetime(2026, 9, 15), 1])
        path = self.save(wb, "izin.xlsx")
        with self.assertRaisesRegex(izin_data_v2.DataError, "Tek park seçin"):
            izin_data_v2.read_izin_v2(path)
        by_park, meta = izin_data_v2.read_izin_v2(path, default_park="DIJITALPARK")
        self.assertEqual(meta["toplam_kisi"], 2)
        self.assertEqual([len(p.gunler) for p in by_park["DIJITALPARK"]], [2, 1])
        results = [{"tc": p.tc, "ad": p.ad, "durum": "zaten_girili", "ok": True,
                    "kaydedildi": False, "flag": []} for p in by_park["DIJITALPARK"]]
        recon = izin_otomasyon.reconcile(by_park["DIJITALPARK"], results,
                                         izin_data_v2.PARKS["DIJITALPARK"])
        self.assertEqual(recon["kayip"], [])
        self.assertFalse(izin_data_v2.PARKS["DIJITALPARK"].onay_dogrulandi)

        incomplete = izin_otomasyon.reconcile(by_park["DIJITALPARK"], results[:1],
                                               izin_data_v2.PARKS["DIJITALPARK"])
        self.assertEqual(len(incomplete["kayip"]), 1)


if __name__ == "__main__":
    unittest.main()
